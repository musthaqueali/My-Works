"""
DamageMech DWG Engine - API 571 Location Mapper
Extracts AutoCAD DWG physical components and maps "Which Damage Mechanism at WHERE"
Generates JSON inspection plan + 5-Sheet Chevron MS Access Excel PCML Database.
"""

import os
import sys
import glob
import subprocess
import argparse
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import ezdxf

import damagemech_rules

ACCORE_PATH = r"C:\Program Files\Autodesk\AutoCAD 2026\accoreconsole.exe"

def convert_dwg_to_dxf(dwg_path: str, temp_dir: str) -> str:
    """Converts native DWG to DXF via accoreconsole."""
    base = os.path.splitext(os.path.basename(dwg_path))[0]
    out_dxf = os.path.join(temp_dir, f"{base}.dxf")
    if os.path.exists(out_dxf) and os.path.getsize(out_dxf) > 1000:
        return out_dxf

    scr_path = os.path.join(temp_dir, f"{base}_export.scr")
    with open(scr_path, "w", encoding="utf-8") as f:
        f.write(f'DXFOUT "{out_dxf}" 16\nQUIT Y\n')

    cmd = [ACCORE_PATH, "/i", dwg_path, "/s", scr_path]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=45)
        if os.path.exists(out_dxf) and os.path.getsize(out_dxf) > 1000:
            return out_dxf
    except Exception as err:
        print(f"[ERROR] DWG conversion failed for {dwg_path}: {err}")
    return None

def parse_dwg_drawing(dxf_path: str) -> dict:
    """Parses DXF entities: title block, fittings, welds, supports, deadlegs, and existing PCMLs."""
    doc = ezdxf.readfile(dxf_path)
    msp = doc.modelspace()

    data = {
        "filename": os.path.basename(dxf_path).replace(".dxf", ".dwg"),
        "drawing_id": os.path.splitext(os.path.basename(dxf_path))[0],
        "plant": "0014t",
        "system": "X001",
        "circuit": "010",
        "sheet_num": "04",
        "full_circuit": "0014t-X001-010",
        "line_tag": "20\"-P-0001-15A",
        "service": "Crude Oil / Hydrocarbon Process",
        "operating_temp": 320.0,
        "operating_press": 150.0,
        "metallurgy": "Carbon Steel (A106-B)",
        "insulation": "NO",
        "pwht": "NO",
        "components": [],
        "pcml_blocks": []
    }

    # Extract block inserts
    for e in msp:
        if e.dxftype() == "INSERT":
            bname = e.dxf.name
            bname_u = bname.upper()
            x = round(e.dxf.insert.x, 2)
            y = round(e.dxf.insert.y, 2)

            # 1. Border Attributes
            if "BORDER" in bname_u or "PLANT INFO" in bname_u:
                for a in getattr(e, "attribs", []):
                    tag = a.dxf.tag.upper()
                    val = a.dxf.text.strip()
                    if tag == "PLANT_NUMBER" and val: data["plant"] = val
                    elif tag == "SYSTEM_NUMBER" and val: data["system"] = val
                    elif tag == "CIRCUIT_NUMBER" and val: data["circuit"] = val
                    elif tag == "SHEET_NUMBER" and val: data["sheet_num"] = val
                    elif tag == "INSULATION" and val: data["insulation"] = val
                    elif tag == "PWHT" and val: data["pwht"] = val
                    elif tag == "LINE_#1" and val and val != "-": data["line_tag"] = val

            # 2. Existing PCML callouts
            elif "PCML" in bname_u:
                blk_info = {"name": bname, "x": x, "y": y, "dm": "85", "dmdisp": "85A", "type": "A"}
                for a in getattr(e, "attribs", []):
                    tag = a.dxf.tag.upper()
                    val = a.dxf.text.strip()
                    if tag == "DM": blk_info["dm"] = val
                    elif tag == "DMDISP": blk_info["dmdisp"] = val
                    elif tag == "TYPE": blk_info["type"] = val
                data["pcml_blocks"].append(blk_info)

            # 3. Elbows
            elif "ELBOW" in bname_u:
                data["components"].append({
                    "name": bname, "type": "ELBOW", "x": x, "y": y, "nps": '20"', "desc": "Butt-Weld Long Radius Elbow"
                })

            # 4. Circumferential Welds
            elif "WELD" in bname_u:
                data["components"].append({
                    "name": bname, "type": "WELD", "x": x, "y": y, "nps": '20"', "desc": "Circumferential Butt Weld & HAZ"
                })

            # 5. Valves
            elif "VALVE" in bname_u:
                data["components"].append({
                    "name": bname, "type": "VALVE", "x": x, "y": y, "nps": '14"', "desc": "Flanged Gate Valve Body"
                })

            # 6. Pipe Supports / Touchpoints
            elif "SUPPORT" in bname_u:
                data["components"].append({
                    "name": bname, "type": "SUPPORT", "x": x, "y": y, "nps": '20"', "desc": "Structural I-Beam Pipe Support Touchpoint"
                })

            # 7. Deadlegs / Bleeders / Drains
            elif any(k in bname_u for k in ["BLEEDER", "DL_", "DRAIN"]):
                data["components"].append({
                    "name": bname, "type": "DEADLEG", "x": x, "y": y, "nps": '2"', "desc": "Stagnant Bleeder / Drain Branch Deadleg"
                })

            # 8. Blind Flanges
            elif "BLIND" in bname_u:
                data["components"].append({
                    "name": bname, "type": "BLIND_FLANGE", "x": x, "y": y, "nps": '14"', "desc": "Threaded / Paddle Blind Flange Cavity"
                })

    data["full_circuit"] = f"{data['plant']}-{data['system']}-{data['circuit']}"

    # Infer fluid stream and metallurgy from line tag if present
    lt = data["line_tag"].upper()
    if "-P-" in lt:
        data["service"] = "Crude Oil / Process Hydrocarbon"
        data["metallurgy"] = "Carbon Steel A106-B"
    elif "-SW-" in lt:
        data["service"] = "Sour Water (H2S / NH3)"
        data["metallurgy"] = "Carbon Steel A106-B"
    elif "-CW-" in lt or "-TW-" in lt:
        data["service"] = "Cooling / Treated Water"
        data["metallurgy"] = "Carbon Steel"

    return data

def run_damagemech_mapping(drawing_data: dict, service_override: str = None, temp_override: float = None) -> dict:
    """Executes API 571 Screening and DWG component location mapping."""
    service = service_override or drawing_data["service"]
    temp = temp_override if temp_override is not None else drawing_data["operating_temp"]
    press = drawing_data["operating_press"]
    mat = drawing_data["metallurgy"]
    ins = drawing_data["insulation"]
    pwht = drawing_data["pwht"]

    # 1. Screen threats for process line
    screened_threats = damagemech_rules.screen_damage_mechanisms(
        service=service,
        temp_f=temp,
        press_psi=press,
        metallurgy=mat,
        insulation=ins,
        pwht=pwht
    )

    # 2. Map threats to DWG components ("Which Damage Mechanism at WHERE")
    mapped_components = damagemech_rules.map_dwg_component_threats(
        components=drawing_data["components"],
        process_threats=screened_threats,
        line_meta=drawing_data
    )

    # 3. Generate PCML / CML inspection table
    cml_points = []
    cml_id_counter = 1

    # First add native DWG PCMLs if present
    for pb in drawing_data["pcml_blocks"]:
        cml_points.append({
            "cml_id": f"CML-{cml_id_counter:02d}",
            "source": "DWG Native Callout",
            "component_id": "COMP-001",
            "sort_dm": pb["dmdisp"],
            "dm_name": "Erosion-Corrosion / Impingement" if "85" in pb["dmdisp"] else "Deadleg / Invert Thinning",
            "location_desc": f"Drawing Callout Point ({pb['dmdisp']})",
            "exact_where": f"Coordinate ({pb['x']}, {pb['y']})",
            "x": pb["x"],
            "y": pb["y"],
            "nps": '20"',
            "ndt_method": "Grid UT (Radial 4-Point)" if "85" in pb["dmdisp"] else "Spot UT (6 o'clock Invert)",
            "susceptibility": "HIGH"
        })
        cml_id_counter += 1

    # Add CML points generated from DWG physical component mapping
    for comp in mapped_components:
        for t in comp["threats"]:
            cml_points.append({
                "cml_id": f"CML-{cml_id_counter:02d}",
                "source": "DamageMech AI Location Engine",
                "component_id": comp["component_id"],
                "sort_dm": t["sort_dm"],
                "dm_name": t["dm_name"],
                "location_desc": f"{comp['component_type']} - {t['exact_location']}",
                "exact_where": f"{comp['desc']} at ({comp['coord_x']}, {comp['coord_y']})",
                "x": comp["coord_x"],
                "y": comp["coord_y"],
                "nps": comp["nps"],
                "ndt_method": t["ndt"],
                "susceptibility": t["susceptibility"]
            })
            cml_id_counter += 1

    result_package = {
        "metadata": drawing_data,
        "screened_threats": screened_threats,
        "mapped_components": mapped_components,
        "cml_points": cml_points,
        "total_components": len(mapped_components),
        "total_threats": len(screened_threats),
        "total_cmls": len(cml_points)
    }

    return result_package

def export_chevron_pcml_excel(mapping_result: dict, output_path: str) -> str:
    """Exports the 5-sheet Chevron MS Access database workbook."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="001489", end_color="001489", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE")
    )

    def write_sheet(ws, headers, rows):
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align

        for r_idx, row in enumerate(rows, start=2):
            ws.append(row)
            for col_idx in range(1, len(row) + 1):
                c = ws.cell(row=r_idx, column=col_idx)
                c.font = data_font
                c.border = thin_border
                c.alignment = left_align if col_idx > 2 else center_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    meta = mapping_result["metadata"]
    circ = meta["full_circuit"]
    iso = meta["drawing_id"]
    sheet_num = meta["sheet_num"]
    line_tag = meta["line_tag"]

    # 1. tblCircuitSheetList
    h1 = ["Original ISO", "Functional Location", "Circuit", "Circuit Sheet", "Sheet Number", "Notes"]
    r1 = [[iso, circ, circ, iso, sheet_num, "DamageMech AI - API 571 Threat Screening & DWG Location Mapped"]]
    ws1 = wb.create_sheet("tblCircuitSheetList")
    write_sheet(ws1, h1, r1)

    # 2. tblCircClass
    h2 = ["Circuit", "PipeClass"]
    r2 = [[circ, "CLASS 2"]]
    ws2 = wb.create_sheet("tblCircClass")
    write_sheet(ws2, h2, r2)

    # 3. tblCircDM (Populated from API 571 Screening!)
    h3 = ["DMKey", "SystemNum", "CircuitNumber", "DMType", "DM", "FailureMod", "Source", "SortDM", "Susceptibility"]
    r3 = []
    dm_key = 1
    for t in mapping_result["screened_threats"]:
        r3.append([
            dm_key,
            "100",
            circ,
            t["category"],
            t["name"],
            t["failure_mode"],
            t["section"],
            t["sort_dm"],
            t["susceptibility"]
        ])
        dm_key += 1
    ws3 = wb.create_sheet("tblCircDM")
    write_sheet(ws3, h3, r3)

    # 4. tblComponents (Populated from DWG components!)
    h4 = ["ComponentID", "Circuit", "LineNumber", "ComponentType", "Description", "NPS", "Schedule", "Material", "X_Coord", "Y_Coord"]
    r4 = []
    for c in mapping_result["mapped_components"]:
        r4.append([
            c["component_id"],
            circ,
            line_tag,
            c["component_type"],
            c["component_name"],
            c["nps"],
            "STD",
            meta["metallurgy"],
            c["coord_x"],
            c["coord_y"]
        ])
    ws4 = wb.create_sheet("tblComponents")
    write_sheet(ws4, h4, r4)

    # 5. tblCMLs (Populated from Location Threat Mapping!)
    h5 = ["CML_ID", "ComponentID", "Circuit", "LineNumber", "DamageMechanism", "SortDM", "ExactLocation", "NDT_Method", "NominalThick", "MinReqThick", "X_Coord", "Y_Coord", "Susceptibility"]
    r5 = []
    for pt in mapping_result["cml_points"]:
        r5.append([
            pt["cml_id"],
            pt["component_id"],
            circ,
            line_tag,
            pt["dm_name"],
            pt["sort_dm"],
            pt["location_desc"],
            pt["ndt_method"],
            0.375,
            0.165,
            pt["x"],
            pt["y"],
            pt["susceptibility"]
        ])
    ws5 = wb.create_sheet("tblCMLs")
    write_sheet(ws5, h5, r5)

    wb.save(output_path)
    print(f"[SUCCESS] Exported 5-Sheet Chevron PCML Excel: {output_path}")
    return output_path

def main():
    parser = argparse.ArgumentParser(description="DamageMech DWG Engine")
    parser.add_argument("--dwg", type=str, default=r"c:\Users\musthaque.mayalankot\OneDrive - Pinnacle\Desktop\DWGs\0014t-X001-010-04.dwg")
    parser.add_argument("--dxf", type=str, default="temp_04.dxf")
    parser.add_argument("--service", type=str, default=None)
    parser.add_argument("--temp", type=float, default=None)
    parser.add_argument("--outJson", type=str, default="damagemech_results.json")
    parser.add_argument("--outXlsx", type=str, default="DamageMech_PCML_Database.xlsx")
    args = parser.parse_args()

    temp_dir = os.path.abspath("scratch_dwg_temp")
    os.makedirs(temp_dir, exist_ok=True)

    dxf_target = args.dxf
    if not os.path.exists(dxf_target):
        if os.path.exists(args.dwg):
            print(f"Converting DWG to DXF: {args.dwg}...")
            dxf_target = convert_dwg_to_dxf(args.dwg, temp_dir)
        else:
            print(f"[ERROR] Neither DXF nor DWG found: {dxf_target}, {args.dwg}")
            sys.exit(1)

    print(f"Parsing DWG/DXF: {dxf_target}...")
    drawing_data = parse_dwg_drawing(dxf_target)
    print(f"Extracted Drawing: Line={drawing_data['line_tag']}, Circuit={drawing_data['full_circuit']}, Components={len(drawing_data['components'])}, PCMLs={len(drawing_data['pcml_blocks'])}")

    print("Running API 571 Screening and Component Location Mapping...")
    results = run_damagemech_mapping(drawing_data, service_override=args.service, temp_override=args.temp)

    # Save JSON
    with open(args.outJson, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)
    print(f"[SUCCESS] Saved JSON Results to: {args.outJson}")

    # Save Excel
    export_chevron_pcml_excel(results, args.outXlsx)
    print(f"Summary: Screened Threats: {results['total_threats']}, Mapped Components: {results['total_components']}, Generated CML Points: {results['total_cmls']}")

if __name__ == "__main__":
    main()
